# -*- coding: utf-8 -*-
"""
ERP_Seed_Data_Sales_Purchasing.xlsx를 sales DB(판매/재고 시트)와
purchase DB(Vendor Contracts/Ratings/Invoice Matching 시트)로 나눠 적재합니다.

이 파일에 있는 구매 관련 시트(Purchase Orders 등)는 이미 purchase DB에 적재된
ERP_Purchasing_Data_Cleaned.xlsx와 행 수가 정확히 일치하는 동일 데이터라 다시
적재하지 않습니다.

실행:
    python scripts/load_sales_purchasing_seed.py /path/to/ERP_Seed_Data_Sales_Purchasing.xlsx
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pymysql

SALES_DB = dict(host="127.0.0.1", user="JangGGo", password="1234", database="sales", charset="utf8mb4")
PURCHASE_DB = dict(host="127.0.0.1", user="JangGGo", password="1234", database="purchase", charset="utf8mb4")

# sales DB로 갈 시트 -> 테이블, 부모 -> 자식 순서
SALES_ORDER = [
    "Customers",
    "Inventory Items",
    "Credit Limits",
    "Customer Contracts",
    "Price Lists",
    "Sales Quotes",
    "Stock Levels",
    "Discounts",
    "Inventory Transfers",
    "Sales Orders",
    "Sales Quote Lines",
    "Order Fulfillment",
    "Sales Order Lines",
    "Fulfillment Lines",
    "Invoices",
    "Sales Reports",
    "Sales Forecasts",
]

# purchase DB로 갈, 이번에 새로 추가된 시트
PURCHASE_NEW_ORDER = [
    "Vendor Contracts",
    "Vendor Ratings",
    "Invoice Matching",
]

SHEET_TO_TABLE = {
    "Customers": "customers",
    "Inventory Items": "inventory_items",
    "Credit Limits": "credit_limits",
    "Customer Contracts": "customer_contracts",
    "Price Lists": "price_lists",
    "Sales Quotes": "sales_quotes",
    "Stock Levels": "stock_levels",
    "Discounts": "discounts",
    "Inventory Transfers": "inventory_transfers",
    "Sales Orders": "sales_orders",
    "Sales Quote Lines": "sales_quote_lines",
    "Order Fulfillment": "order_fulfillment",
    "Sales Order Lines": "sales_order_lines",
    "Fulfillment Lines": "fulfillment_lines",
    "Invoices": "invoices",
    "Sales Reports": "sales_reports",
    "Sales Forecasts": "sales_forecasts",
    "Vendor Contracts": "vendor_contracts",
    "Vendor Ratings": "vendor_ratings",
    "Invoice Matching": "invoice_matching",
}


def header_to_column(header: str) -> str:
    return header.strip().lower()


def _normalize_value(value):
    """엑셀에서 'TRUE'/'FALSE' 문자열로 들어오는 값을 MySQL BOOLEAN에 맞게 변환합니다."""
    if isinstance(value, str):
        if value.strip().upper() == "TRUE":
            return 1
        if value.strip().upper() == "FALSE":
            return 0
    return value


def load_sheet(cursor, workbook, sheet_name: str) -> int:
    table = SHEET_TO_TABLE[sheet_name]
    ws = workbook[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    header = [header_to_column(h) for h in rows[0]]
    data_rows = [tuple(_normalize_value(v) for v in row) for row in rows[1:]]
    if not data_rows:
        return 0

    placeholders = ", ".join(["%s"] * len(header))
    columns = ", ".join(f"`{c}`" for c in header)
    sql = f"INSERT INTO `{table}` ({columns}) VALUES ({placeholders})"
    cursor.executemany(sql, data_rows)
    return len(data_rows)


def load_into(db_config: dict, workbook, sheet_order: list[str]) -> None:
    connection = pymysql.connect(**db_config)
    try:
        with connection.cursor() as cursor:
            cursor.execute("SET FOREIGN_KEY_CHECKS = 0")
            for sheet_name in reversed(sheet_order):
                cursor.execute(f"DELETE FROM `{SHEET_TO_TABLE[sheet_name]}`")
            cursor.execute("SET FOREIGN_KEY_CHECKS = 1")

            total = 0
            for sheet_name in sheet_order:
                count = load_sheet(cursor, workbook, sheet_name)
                print(f"[{db_config['database']}] {sheet_name} -> {SHEET_TO_TABLE[sheet_name]}: {count}행")
                total += count
        connection.commit()
        print(f"[{db_config['database']}] 총 {total}행 적재 완료.\n")
    finally:
        connection.close()


def main() -> None:
    if len(sys.argv) < 2:
        print("사용법: python scripts/load_sales_purchasing_seed.py <엑셀_파일_경로>")
        sys.exit(1)

    excel_path = Path(sys.argv[1])
    workbook = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)

    load_into(SALES_DB, workbook, SALES_ORDER)
    load_into(PURCHASE_DB, workbook, PURCHASE_NEW_ORDER)


if __name__ == "__main__":
    main()
